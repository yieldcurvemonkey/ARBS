"""Tests for the Citi Velocity intraday USD-SOFR rateslib curve package.

Hermetic unit tests build a synthetic ``CVTSHIST``-layout workbook in ``tmp_path``
so they run in the fast gate with no network / external files.  A single
``integration``-marked test additionally exercises the real ``db.xlsx`` export
when it is present, and self-skips otherwise.
"""

from __future__ import annotations

import datetime

import openpyxl
import pandas as pd
import pytest
import rateslib as rl

from MDP.IRSwaps.CITI_VELOCITY_INTRADAY import (
    CitiVelocityIntradayFetcher,
    CitiVelocityWorkbook,
    CURVE_TENOR_ORDER,
    DEFAULT_DB_PATH,
    build_rl_usd_sofr_intraday_curve,
    load_intraday_par_rates,
    parse_period_from_sheet_name,
    parse_tenor_from_header,
)
from MDP.IRSwaps.CITI_VELOCITY_INTRADAY.rl_usd_sofr_intraday_builder import (
    par_reprice_errors_bp,
    spot_date,
)
from MDP.IRSwaps.SDR_INTRADAY.rl_curve_utils.stir_curve_building_utils import RLCurveBase


# Representative short→long grid + plausible, monotone-ish par rates (percent).
FIXTURE_TENORS = ["1D", "1W", "1M", "3M", "6M", "1Y", "2Y", "5Y", "10Y", "30Y"]
FIXTURE_RATES_A = [4.30, 4.31, 4.28, 4.20, 4.05, 3.85, 3.60, 3.55, 3.70, 3.85]
FIXTURE_RATES_B = [4.31, 4.32, 4.29, 4.21, 4.06, 3.86, 3.61, 3.56, 3.71, 3.86]


def _write_citi_sheet(ws, start_int, end_int, timestamps, rows, tenors=FIXTURE_TENORS):
    """Populate one worksheet with the Citi CVTSHIST layout."""
    ws.cell(row=1, column=1, value=start_int)
    ws.cell(row=2, column=1, value=end_int)
    # row 3 blank; row 4 formula
    ws.cell(row=4, column=1, value='=CVTSHIST("RATES.OIS.USD_SOFR.PAR.1D,...","MI01",,"x","y","CLOSE")')
    # row 5 header
    ws.cell(row=5, column=1, value="Date")
    for j, t in enumerate(tenors, start=2):
        ws.cell(row=5, column=j, value=f"RATES.OIS.USD_SOFR.PAR.{t} - CLOSE")
    # row 6+ data (newest first, like the real export)
    for i, (ts, vals) in enumerate(zip(timestamps, rows), start=6):
        ws.cell(row=i, column=1, value=ts)
        for j, v in enumerate(vals, start=2):
            ws.cell(row=i, column=j, value=v)


@pytest.fixture()
def citi_xlsx(tmp_path):
    """A two-sheet workbook: one populated week + one empty template week."""
    path = tmp_path / "citi_fixture.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "202607130001-202607171159"
    # two snapshots, newest first (11:58 then 11:57)
    ts_new = datetime.datetime(2026, 7, 17, 11, 58)
    ts_old = datetime.datetime(2026, 7, 17, 11, 57)
    _write_citi_sheet(
        ws,
        202607130001,
        202607171159,
        [ts_new, ts_old],
        [FIXTURE_RATES_A, FIXTURE_RATES_B],
    )
    # empty template sheet (only the two boundary timestamps)
    ws2 = wb.create_sheet("202607200001-202607241159")
    ws2.cell(row=1, column=1, value=202607200001)
    ws2.cell(row=2, column=1, value=202607241159)
    wb.save(path)
    wb.close()
    return str(path)


# --------------------------------------------------------------------------- #
# header / metadata parsing
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    "header,expected",
    [
        ("RATES.OIS.USD_SOFR.PAR.5Y - CLOSE", "5Y"),
        ("RATES.OIS.USD_SOFR.PAR.15M - CLOSE", "15M"),
        ("RATES.OIS.USD_SOFR.PAR.1D - CLOSE", "1D"),
        ("Date", None),
        (None, None),
    ],
)
def test_parse_tenor_from_header(header, expected):
    assert parse_tenor_from_header(header) == expected


def test_parse_period_from_sheet_name():
    start, end = parse_period_from_sheet_name("202607130001-202607171159")
    assert start == datetime.datetime(2026, 7, 13, 0, 1)
    assert end == datetime.datetime(2026, 7, 17, 11, 59)
    assert parse_period_from_sheet_name("Sheet1") is None
    assert parse_period_from_sheet_name("not-a-period") is None


def test_curve_tenor_order_has_44_entries():
    assert len(CURVE_TENOR_ORDER) == 44
    assert CURVE_TENOR_ORDER[0] == "1D"
    assert CURVE_TENOR_ORDER[-1] == "50Y"


# --------------------------------------------------------------------------- #
# loader
# --------------------------------------------------------------------------- #
def test_loader_parses_populated_skips_empty(citi_xlsx):
    frames = load_intraday_par_rates(citi_xlsx, concat=False)
    assert set(frames) == {"202607130001-202607171159"}  # empty template skipped
    df = frames["202607130001-202607171159"]
    assert list(df.columns) == FIXTURE_TENORS  # ordered per CURVE_TENOR_ORDER
    assert df.shape == (2, len(FIXTURE_TENORS))
    # ascending order (oldest first after sort)
    assert list(df.index) == [
        datetime.datetime(2026, 7, 17, 11, 57),
        datetime.datetime(2026, 7, 17, 11, 58),
    ]
    # values line up with the newest row
    assert df.loc[datetime.datetime(2026, 7, 17, 11, 58), "10Y"] == pytest.approx(3.70)


def test_loader_concat_and_workbook(citi_xlsx):
    combined = load_intraday_par_rates(citi_xlsx, concat=True)
    assert isinstance(combined, pd.DataFrame)
    assert combined.shape == (2, len(FIXTURE_TENORS))

    wb = CitiVelocityWorkbook(citi_xlsx)
    assert len(wb.sheet_names()) == 2
    assert wb.populated_sheets() == ["202607130001-202607171159"]
    assert wb.frame().shape == (2, len(FIXTURE_TENORS))
    periods = wb.sheet_periods()
    assert set(periods) == {"202607130001-202607171159", "202607200001-202607241159"}


def test_loader_rejects_non_xlsx(tmp_path):
    p = tmp_path / "legacy.xls"
    p.write_bytes(b"\xd0\xcf\x11\xe0")  # OLE2 magic
    with pytest.raises(ValueError, match="xlsx"):
        load_intraday_par_rates(str(p))


def test_snapshot_selection_methods(citi_xlsx):
    wb = CitiVelocityWorkbook(citi_xlsx)
    latest = wb.snapshot()
    assert latest.name == datetime.datetime(2026, 7, 17, 11, 58)
    assert latest["1D"] == pytest.approx(4.30)

    asof = wb.snapshot(datetime.datetime(2026, 7, 17, 11, 57, 30), method="asof")
    assert asof.name == datetime.datetime(2026, 7, 17, 11, 57)

    nearest = wb.snapshot(datetime.datetime(2026, 7, 17, 11, 57, 40), method="nearest")
    assert nearest.name == datetime.datetime(2026, 7, 17, 11, 58)

    exact = wb.snapshot(datetime.datetime(2026, 7, 17, 11, 57), method="exact")
    assert exact.name == datetime.datetime(2026, 7, 17, 11, 57)


# --------------------------------------------------------------------------- #
# builder
# --------------------------------------------------------------------------- #
def _snapshot_series(rates=FIXTURE_RATES_A, tenors=FIXTURE_TENORS, ts=datetime.datetime(2026, 7, 17, 11, 58)):
    return pd.Series(dict(zip(tenors, rates)), name=ts)


def test_spot_date_is_t_plus_2_good_business_days():
    # 2026-07-17 is a Friday; +2 good business days -> Tue 2026-07-21
    assert spot_date(datetime.date(2026, 7, 17)) == rl.dt(2026, 7, 21)


def test_spot_date_rolls_weekend_and_holiday_to_prior_business_day():
    # 2026-01-16 Fri; 17-18 weekend; 19 MLK holiday -> all anchor to Fri 01-16,
    # so spot = +2 good business days past MLK week = Wed 2026-01-21.
    friday = spot_date(datetime.date(2026, 1, 16))
    assert friday == rl.dt(2026, 1, 21)
    assert spot_date(datetime.date(2026, 1, 17)) == friday          # Saturday
    assert spot_date(datetime.date(2026, 1, 18)) == friday          # Sunday
    assert spot_date(datetime.datetime(2026, 1, 19, 15, 0)) == friday  # MLK Monday


def test_builder_holiday_ref_anchors_prior_business_day():
    # A snapshot dated on MLK Monday must NOT crash; it anchors to Fri 2026-01-16.
    s = _snapshot_series(ts=datetime.datetime(2026, 1, 19, 15, 0))
    rlc = build_rl_usd_sofr_intraday_curve(par_rates=s, ref_date=s.name)
    assert rlc.rl_pricing_curve_solver.result["status"] == "SUCCESS"
    # every calibrating swap starts at the Friday-anchored spot date
    eff = list(rlc.rl_pricing_curve_instruments.values())[0].leg1.schedule.effective
    assert eff == rl.dt(2026, 1, 21)
    assert par_reprice_errors_bp(rlc).abs().max() < 1e-2


def test_builder_min_tenors_guard():
    # default floor (4) rejects a 3-tenor snapshot ...
    with pytest.raises(ValueError, match="min_tenors"):
        build_rl_usd_sofr_intraday_curve(
            par_rates={"1D": 4.3, "1Y": 3.9, "10Y": 4.0}, ref_date=datetime.date(2026, 7, 15)
        )
    # ... but an explicit override lets a tiny curve through.
    rlc = build_rl_usd_sofr_intraday_curve(
        par_rates={"1D": 4.3, "10Y": 4.0}, ref_date=datetime.date(2026, 7, 15), min_tenors=1
    )
    assert rlc.rl_pricing_curve_solver.result["status"] == "SUCCESS"


def test_builder_solves_and_reprices():
    s = _snapshot_series()
    rlc = build_rl_usd_sofr_intraday_curve(par_rates=s, ref_date=s.name, timestamp=s.name)
    assert isinstance(rlc, RLCurveBase)
    assert rlc.rl_pricing_curve_solver.result["status"] == "SUCCESS"
    assert rlc.rl_risk_curve is None
    # one calibrating instrument per supplied tenor
    assert set(rlc.rl_pricing_curve_instruments) == set(FIXTURE_TENORS)

    errs = par_reprice_errors_bp(rlc)
    assert errs.abs().max() < 1e-3  # sub-1e-3 bp
    # every supplied tenor re-priced
    assert set(errs.index) == set(FIXTURE_TENORS)


def test_builder_discount_factors_monotone_decreasing():
    s = _snapshot_series()
    rlc = build_rl_usd_sofr_intraday_curve(par_rates=s, ref_date=s.name)
    curve = rlc.rl_pricing_curve
    mats = [irs.leg1.schedule.termination for irs in rlc.rl_pricing_curve_instruments.values()]
    mats = sorted(mats)
    dfs = [float(curve[d]) for d in mats]
    assert all(dfs[i] >= dfs[i + 1] for i in range(len(dfs) - 1))
    assert dfs[0] < 1.0 and dfs[-1] > 0.0


def test_builder_spline_variant_solves():
    s = _snapshot_series()
    rlc = build_rl_usd_sofr_intraday_curve(par_rates=s, ref_date=s.name, spline_start_tenor="2Y")
    assert rlc.rl_pricing_curve_solver.result["status"] == "SUCCESS"
    assert par_reprice_errors_bp(rlc).abs().max() < 1e-2


def test_builder_empty_rates_raises():
    with pytest.raises(ValueError, match="no usable par rates"):
        build_rl_usd_sofr_intraday_curve(par_rates={}, ref_date=datetime.date(2026, 7, 17))


def test_builder_drops_nan_rates():
    s = _snapshot_series()
    s = s.copy()
    s["5Y"] = float("nan")
    rlc = build_rl_usd_sofr_intraday_curve(par_rates=s, ref_date=s.name)
    assert "5Y" not in rlc.rl_pricing_curve_instruments
    assert len(rlc.rl_pricing_curve_instruments) == len(FIXTURE_TENORS) - 1


def test_builder_bad_spline_tenor_raises():
    s = _snapshot_series()
    with pytest.raises(ValueError, match="spline_start_tenor"):
        build_rl_usd_sofr_intraday_curve(par_rates=s, ref_date=s.name, spline_start_tenor="7Y")


# --------------------------------------------------------------------------- #
# fetcher
# --------------------------------------------------------------------------- #
def test_fetcher_build_curve_and_latest(citi_xlsx):
    f = CitiVelocityIntradayFetcher(citi_xlsx)
    assert len(f.timestamps()) == 2
    rlc = f.build_curve()
    assert rlc.timestamp == datetime.datetime(2026, 7, 17, 11, 58)
    assert rlc.rl_pricing_curve_solver.result["status"] == "SUCCESS"
    # latest == build_curve(None)
    assert f.latest().timestamp == rlc.timestamp


def test_fetcher_eris_style_return_shapes(citi_xlsx):
    f = CitiVelocityIntradayFetcher(citi_xlsx)

    curve, ts = f.fetch_intraday_discount_curve("2026-07-17 11:57:40")
    assert isinstance(curve, rl.Curve)
    assert ts == datetime.datetime(2026, 7, 17, 11, 58)

    curve_only = f.fetch_intraday_discount_curve(return_intraday_timestamp=False)
    assert isinstance(curve_only, rl.Curve)

    rlc, ts2 = f.fetch_intraday_discount_curve(return_rl_curve_base=True)
    assert isinstance(rlc, RLCurveBase)

    par = f.fetch_intraday_discount_curve(return_df=True)
    assert isinstance(par, pd.Series)
    assert par["1D"] == pytest.approx(4.30)


def test_fetcher_cache_returns_same_object(citi_xlsx):
    f = CitiVelocityIntradayFetcher(citi_xlsx)
    a = f.build_curve("2026-07-17 11:58", method="exact")
    b = f.build_curve("2026-07-17 11:58", method="exact")
    assert a is b  # cache hit


def test_fetcher_iter_curves(citi_xlsx):
    f = CitiVelocityIntradayFetcher(citi_xlsx)
    # no freq: every snapshot in range
    out = list(f.iter_curves())
    assert len(out) == 2
    assert all(isinstance(rlc, RLCurveBase) for _, rlc in out)
    # freq grid snaps to nearest, duplicates collapsed -> <= 2 distinct
    out2 = list(f.iter_curves("2026-07-17 11:00", "2026-07-17 12:00", freq="1min"))
    assert 1 <= len(out2) <= 2


def test_fetcher_iter_curves_asof_method_does_not_crash(citi_xlsx):
    # "asof" is a documented method; the freq-grid path must not forward it raw
    # to Index.get_indexer (which only accepts pad/backfill/nearest).
    f = CitiVelocityIntradayFetcher(citi_xlsx)
    out = list(f.iter_curves("2026-07-17 11:00", "2026-07-17 12:00", freq="1min", method="asof"))
    assert all(isinstance(rlc, RLCurveBase) for _, rlc in out)


def test_workbook_metadata_rejects_non_xlsx(tmp_path):
    # sheet_names/sheet_periods must reject .xls with the same friendly ValueError
    # as the data-loading path (not a raw openpyxl exception).
    p = tmp_path / "legacy.xls"
    p.write_bytes(b"\xd0\xcf\x11\xe0")
    wb = CitiVelocityWorkbook(str(p))
    with pytest.raises(ValueError, match="xlsx"):
        wb.sheet_periods()
    with pytest.raises(ValueError, match="xlsx"):
        wb.sheet_names()


def test_fetcher_custom_curve_id(citi_xlsx):
    f = CitiVelocityIntradayFetcher(citi_xlsx, curve_id="USD-SOFR-CITI")
    rlc = f.build_curve()
    assert rlc.rl_pricing_curve.id == "USD-SOFR-CITI"


# --------------------------------------------------------------------------- #
# real-data integration (skips when db.xlsx absent)
# --------------------------------------------------------------------------- #
@pytest.mark.integration
def test_real_db_xlsx_builds_and_reprices():
    import os

    if not os.path.exists(DEFAULT_DB_PATH):
        pytest.skip(f"Citi Velocity workbook not present at {DEFAULT_DB_PATH}")

    f = CitiVelocityIntradayFetcher(DEFAULT_DB_PATH)
    populated = f.populated_sheets()
    assert populated, "expected at least one populated sheet"

    rlc = f.build_curve()  # latest snapshot
    assert len(rlc.rl_pricing_curve_instruments) == 44  # full tenor grid
    assert rlc.rl_pricing_curve_solver.result["status"] == "SUCCESS"
    errs = par_reprice_errors_bp(rlc)
    assert errs.abs().max() < 0.01  # sub-0.01 bp against real data


# --------------------------------------------------------------------------- #
# IRSwapsMDP integration (source='citivelo'); needs SOFR fixings -> network
# --------------------------------------------------------------------------- #
@pytest.mark.integration
@pytest.mark.network
def test_irswapsmdp_citivelo_source(citi_xlsx):
    from MDP.IRSwaps.IRSwapsMDP import IRSwapsMDP
    from Query.IRSwaps.backends.rateslib.RLIRSwapCurve import RLIRSwapCurve

    mdp = IRSwapsMDP(source="citivelo")

    # 'live' -> latest snapshot in the (fixture) workbook
    curve = mdp.get_pricer(
        {"curve_name": "USD-SOFR-1D", "timestamp": "live", "workbook_path": citi_xlsx}
    )
    assert isinstance(curve, RLIRSwapCurve)
    assert curve.id() == "USD-SOFR-1D"
    assert str(curve.meta()["id"]).startswith("CITIVELO-USD-SOFR-1D-")
    ts = curve.meta()["timestamp"]
    assert ts.date() == datetime.date(2026, 7, 17)
    # anchor node is the (business-day-rolled) snapshot date
    assert curve.reference_date().date() == datetime.date(2026, 7, 17)

    # the wrapped curve prices a spot-starting swap sensibly
    swap = curve.build_irswap(fwd="0D", tenor="10Y")
    assert 0.0 < curve.fair_rate(swap) * 100.0 < 15.0

    # explicit datetime request -> nearest snapshot
    curve2 = mdp.get_pricer(
        {"curve_name": "USD-SOFR-1D", "timestamp": datetime.datetime(2026, 7, 17, 11, 57, 40),
         "workbook_path": citi_xlsx}
    )
    assert curve2.meta()["timestamp"] == datetime.datetime(2026, 7, 17, 11, 58)

    # non-SOFR curve name trips the guard
    with pytest.raises(AssertionError):
        mdp.get_pricer({"curve_name": "EUR-ESTR", "timestamp": "live", "workbook_path": citi_xlsx})
