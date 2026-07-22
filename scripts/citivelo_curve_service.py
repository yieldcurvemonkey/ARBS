"""Warm the CurveStore with Citi Velocity intraday USD-SOFR curves.

Reads the (now fully-populated, ~365 MB) Citi Velocity ``CVTSHIST`` workbook of
USD SOFR OIS **par** rates, solves one rateslib discount curve per 1-minute
snapshot, and writes them into the local CurveStore — mirroring the optimize
logic used by ``scripts/stirf_curve_service.py`` for the intraday BARCHART_STIRF
curves (day-bucketed work, skip-already-populated, process-pool calibration,
BackfillProgress-style ETA/RSS heartbeat, direct ``store.write_day``).

Design decisions (see the module docstring notes):

* **Dedicated asset name** ``USD-SOFR-1D-CITIVELO``.  The store partitions only
  by ``asset={curve_name}`` and reads do NOT filter by ``source_variant``, so
  the shared ``USD-SOFR-1D`` asset (which already holds ERIS EOD history) must
  not be reused — that would mix on read / clobber ERIS on overwrite.
* **Local-only writes.**  ``ARBS_SUPABASE_ENABLED=0`` is exported *before* any
  ``Caching`` import so ``write_day`` never background-pushes ~1M rows to the
  remote prod Supabase (which is enabled by default via baked-in credentials).
* **Calendar-date partitions.**  Each snapshot is bucketed by its ET calendar
  date (not the CME 17:00-CT trading-date roll), so a whole day lands in one
  partition — one idempotent ``write_day(overwrite=True)`` per day, safe for the
  process pool, and the natural "nearest snapshot within a day" access pattern.
* **Solver reuse.**  Within a day the node dates are constant, so one
  ``rl.Solver`` is built per NaN-tenor signature and re-solved per minute via
  ``solver.s = rates; solver.iterate()`` (~2.7x faster than cold rebuilds,
  reprices to <0.002 bp).

Usage::

    conda run -n stir python scripts/citivelo_curve_service.py warm
    conda run -n stir python scripts/citivelo_curve_service.py warm --start-date 2026-01-01 --verify
    conda run -n stir python scripts/citivelo_curve_service.py status
"""

from __future__ import annotations

# ── MUST come before any Caching import (also re-runs in spawned workers) ──
import os

os.environ.setdefault("ARBS_SUPABASE_ENABLED", "0")

import argparse
import contextlib
import datetime
import gc
import io
import logging
import sys
import time
from concurrent.futures import ProcessPoolExecutor, as_completed
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterable, Optional

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytz
import rateslib as rl

# Repo modules (import after the env toggle above).
from MDP.IRSwaps.CITI_VELOCITY_INTRADAY.CitiVelocityIntradayFetcher import DEFAULT_DB_PATH
from MDP.IRSwaps.CITI_VELOCITY_INTRADAY.citi_velocity_loader import parse_sheet
from MDP.IRSwaps.CITI_VELOCITY_INTRADAY.rl_usd_sofr_intraday_builder import (
    _prev_business_day,
    _tenor_sort_key,
)

# ── constants ──
ASSET_NAME = "USD-SOFR-1D-CITIVELO"
REFERENCE_KEY = "USD-SOFR-1D"          # drives act360/nyc/mf on reconstruct
SOURCE_VARIANT = "CITIVELO"
DEFAULT_INTERPOLATION = "log_linear"
DEFAULT_CONVENTION = "act360"
DEFAULT_CALENDAR = "nyc"
DEFAULT_MODIFIER = "mf"
DEFAULT_SPOT_LAG = 2
MIN_TENORS = 4
SOURCE_TZ = "America/New_York"          # naive workbook wall-clock tz

_ET = pytz.timezone(SOURCE_TZ)
_UTC = pytz.UTC
_CHI = pytz.timezone("America/Chicago")

LOGGER = logging.getLogger("citivelo_curve_service")


# ─────────────────────────────────────────────────────────────────────────────
# Phase 1 — extract the workbook into per-calendar-date par-rate parquet files
# ─────────────────────────────────────────────────────────────────────────────
def _default_work_dir() -> Path:
    base = os.environ.get("ARBS_CACHE_DIR")
    root = Path(base) if base else Path.home() / "AppData" / "Local" / "ARBS" / "Cache"
    return root / "citivelo_par_extract"


def extract_par_rates(
    workbook: str,
    work_dir: Path,
    *,
    force: bool = False,
    logger: logging.Logger = LOGGER,
) -> list[datetime.date]:
    """Stream the workbook → one ``{date}.parquet`` of (timestamp, tenors) per day.

    Idempotent: an existing day file is left untouched unless ``force``.  Returns
    the sorted list of calendar dates present in ``work_dir`` afterwards.
    """
    import openpyxl

    work_dir.mkdir(parents=True, exist_ok=True)
    t0 = time.time()
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        populated = [sn for sn in wb.sheetnames if (wb[sn].max_column or 0) > 1]
        logger.info("extract: %d populated sheets in %s", len(populated), workbook)
        n_written = n_rows = 0
        for i, sn in enumerate(populated, 1):
            df = parse_sheet(wb[sn])
            if df.empty:
                continue
            for d, sub in df.groupby(df.index.date):
                out = work_dir / f"{d.isoformat()}.parquet"
                if out.exists() and not force:
                    continue
                arrays = [pa.array(sub.index.values.astype("datetime64[us]"))]
                names = ["timestamp"]
                for c in sub.columns:
                    arrays.append(pa.array(sub[c].to_numpy(), type=pa.float64()))
                    names.append(c)
                table = pa.Table.from_arrays(arrays, names=names)
                tmp = out.with_suffix(".parquet.tmp")
                pq.write_table(table, tmp, compression="zstd")
                os.replace(tmp, out)
                n_written += 1
                n_rows += len(sub)
            if i % 20 == 0 or i == len(populated):
                logger.info(
                    "extract: sheet %d/%d  days_written=%d rows=%d elapsed=%.1fs",
                    i, len(populated), n_written, n_rows, time.time() - t0,
                )
    finally:
        wb.close()

    dates = sorted(
        datetime.date.fromisoformat(p.stem)
        for p in work_dir.glob("*.parquet")
    )
    logger.info("extract: %d day files present in %s (%.1fs)", len(dates), work_dir, time.time() - t0)
    return dates


# ─────────────────────────────────────────────────────────────────────────────
# Phase 2 — build curves per day + write CurveSnapshots (process pool workers)
# ─────────────────────────────────────────────────────────────────────────────
@dataclass
class DayStat:
    date: datetime.date
    n_curves: int = 0
    n_fallback: int = 0
    n_skipped_rows: int = 0
    elapsed: float = 0.0
    error: Optional[str] = None


# per-process singletons (populated by the pool initializer)
_WORKER: dict[str, Any] = {"store": None, "params": None}


def _init_worker(base_dir: Optional[str]) -> None:
    # env already forced local-only at module import; build the store once/process.
    from Caching.curve_store import CurveStore

    _WORKER["store"] = CurveStore(base_dir=Path(base_dir)) if base_dir else CurveStore.default()


def _build_day_snapshots(date: datetime.date, par_path: str, params: dict) -> tuple[list, DayStat]:
    """Build all CurveSnapshots for one calendar day (no store I/O here)."""
    from Caching.curve_store import CurveSnapshot

    interpolation = params["interpolation"]
    convention = params["convention"]
    calendar = params["calendar"]
    modifier = params["modifier"]
    spot_lag = params["spot_lag"]

    df = pd.read_parquet(par_path).set_index("timestamp").sort_index()
    stat = DayStat(date=date)

    ref = _prev_business_day(rl.dt(date.year, date.month, date.day), calendar)
    cal = rl.get_calendar(calendar)
    spot = cal.add_bus_days(ref, spot_lag, True)

    snapshots: list = []
    # group minutes by their non-NaN tenor signature so the node set is constant
    notna = df.notna()
    sig = notna.apply(lambda r: tuple(df.columns[r.values]), axis=1)
    for signature, idx in sig.groupby(sig).groups.items():
        tenors = sorted(signature, key=_tenor_sort_key)
        if len(tenors) < MIN_TENORS:
            stat.n_skipped_rows += len(idx)
            continue
        protos = [
            (t, rl.IRS(effective=spot, termination=t, spec="usd_irs", curves="c", fixed_rate=2.0))
            for t in tenors
        ]
        protos.sort(key=lambda p: p[1].leg1.schedule.termination)
        ordered = [p[0] for p in protos]
        node_dates = [p[1].leg1.schedule.termination for p in protos]
        nodes_template = {ref: 1.0, **{d: 1.0 for d in node_dates}}

        group_df = df.loc[idx, ordered].sort_index()
        rows = group_df.to_numpy(dtype="float64")
        timestamps = group_df.index.to_pydatetime()

        instrs = [
            rl.IRS(effective=spot, termination=t, spec="usd_irs", curves="c", fixed_rate=float(rows[0][j]))
            for j, t in enumerate(ordered)
        ]
        curve = rl.Curve(nodes=dict(nodes_template), id="c", convention=convention,
                         calendar=calendar, modifier=modifier, interpolation=interpolation)
        solver = rl.Solver(curves=[curve], instruments=instrs, s=list(rows[0]), id="c",
                           func_tol=1e-9, conv_tol=1e-10)

        for k in range(len(rows)):
            rates = list(rows[k])
            solver.s = rates
            solver.iterate()
            if solver.result.get("status") != "SUCCESS":
                # cold fallback: fresh curve + solver for this minute
                curve = rl.Curve(nodes=dict(nodes_template), id="c", convention=convention,
                                 calendar=calendar, modifier=modifier, interpolation=interpolation)
                solver = rl.Solver(curves=[curve], instruments=instrs, s=rates, id="c",
                                   func_tol=1e-9, conv_tol=1e-10)
                stat.n_fallback += 1
                if solver.result.get("status") != "SUCCESS":
                    stat.n_skipped_rows += 1
                    continue

            raw = curve.nodes._nodes if hasattr(curve.nodes, "_nodes") else dict(curve.nodes)
            nd_sorted = sorted(raw.keys())
            ts_naive = timestamps[k]
            ts_utc = _ET.localize(ts_naive).astimezone(_UTC).replace(microsecond=0)
            snapshots.append(
                CurveSnapshot(
                    timestamp_utc=ts_utc,
                    timestamp_local=ts_utc.astimezone(_CHI),
                    trading_date=date,                       # ET calendar-date partition
                    session_minute=int(ts_naive.hour * 60 + ts_naive.minute),  # minute-of-day (ET)
                    curve_name=ASSET_NAME,
                    cfg_hash="",
                    reference_key=REFERENCE_KEY,
                    interpolation=interpolation,
                    source_variant=SOURCE_VARIANT,
                    node_dates=[d.date() if hasattr(d, "date") else d for d in nd_sorted],
                    discount_factors=[float(raw[d]) for d in nd_sorted],
                )
            )

    stat.n_curves = len(snapshots)
    return snapshots, stat


def _warm_one_day(task: tuple) -> DayStat:
    date, par_path, params = task
    t0 = time.time()
    try:
        # rateslib prints a "SUCCESS: func_tol reached ..." line per solve; with
        # ~1M solves across many workers that both bloats the log and serializes
        # on the shared stdout. Swallow it — workers communicate via the return value.
        with contextlib.redirect_stdout(io.StringIO()):
            snapshots, stat = _build_day_snapshots(date, par_path, params)
        if snapshots:
            snapshots.sort(key=lambda s: s.timestamp_utc)
            _WORKER["store"].write_day(ASSET_NAME, date, snapshots, overwrite=True)
        stat.elapsed = time.time() - t0
        return stat
    except Exception as exc:  # isolate one bad day
        return DayStat(date=date, elapsed=time.time() - t0, error=f"{type(exc).__name__}: {exc}")
    finally:
        gc.collect()


# ─────────────────────────────────────────────────────────────────────────────
# progress
# ─────────────────────────────────────────────────────────────────────────────
def _rss_mb() -> float:
    try:
        import psutil

        return psutil.Process().memory_info().rss / 1e6
    except Exception:
        return 0.0


@dataclass
class Progress:
    total_days: int
    started: float = field(default_factory=time.time)
    days_done: int = 0
    curves: int = 0
    fallbacks: int = 0
    errors: int = 0

    def record(self, stat: DayStat) -> None:
        self.days_done += 1
        self.curves += stat.n_curves
        self.fallbacks += stat.n_fallback
        if stat.error:
            self.errors += 1

    def heartbeat(self) -> str:
        elapsed = time.time() - self.started
        pct = 100.0 * self.days_done / self.total_days if self.total_days else 100.0
        rate = self.days_done / (elapsed / 60.0) if elapsed > 0 else 0.0
        remaining = self.total_days - self.days_done
        eta = (remaining / rate * 60.0) if rate > 0 else 0.0
        cps = self.curves / elapsed if elapsed > 0 else 0.0
        return (
            f"day {self.days_done}/{self.total_days} ({pct:.1f}%) "
            f"curves={self.curves:,} ({cps:.0f}/s) fallbacks={self.fallbacks} errors={self.errors} "
            f"elapsed={elapsed/60:.1f}m eta={eta/60:.1f}m rate={rate:.1f} days/min mem={_rss_mb():.0f}MB"
        )


# ─────────────────────────────────────────────────────────────────────────────
# orchestration
# ─────────────────────────────────────────────────────────────────────────────
def _select_dates(
    all_dates: list[datetime.date],
    *,
    start: Optional[datetime.date],
    end: Optional[datetime.date],
    limit: Optional[int],
) -> list[datetime.date]:
    out = [d for d in all_dates if (start is None or d >= start) and (end is None or d <= end)]
    if limit:
        out = out[:limit]
    return out


def _pending_dates(store: Any, dates: list[datetime.date], *, force: bool) -> list[datetime.date]:
    if force:
        return list(dates)
    return [d for d in dates if not store.has_day(ASSET_NAME, d)]


def run_warm(args: argparse.Namespace, logger: logging.Logger) -> int:
    from Caching.curve_store import CurveStore

    work_dir = Path(args.work_dir) if args.work_dir else _default_work_dir()
    base_dir = args.base_dir

    # Phase 1: extract
    if args.skip_extract:
        all_dates = sorted(
            datetime.date.fromisoformat(p.stem) for p in work_dir.glob("*.parquet")
        )
        logger.info("extract skipped: reusing %d day files in %s", len(all_dates), work_dir)
    else:
        all_dates = extract_par_rates(args.workbook, work_dir, force=args.force_extract, logger=logger)

    start = datetime.date.fromisoformat(args.start_date) if args.start_date else None
    end = datetime.date.fromisoformat(args.end_date) if args.end_date else None
    selected = _select_dates(all_dates, start=start, end=end, limit=args.limit_days)

    store = CurveStore(base_dir=Path(base_dir)) if base_dir else CurveStore.default()
    logger.info("curve store base_dir: %s  (asset=%s)", store._base_dir, ASSET_NAME)
    from Caching.supabase_engine import SUPABASE_ENABLED
    logger.info("SUPABASE_ENABLED=%s (writes are %s)", SUPABASE_ENABLED,
                "LOCAL-ONLY" if not SUPABASE_ENABLED else "PUSHING REMOTE -- ABORT")
    if SUPABASE_ENABLED:
        logger.error("Refusing to run: remote Supabase push is enabled. Export ARBS_SUPABASE_ENABLED=0.")
        return 2

    pending = _pending_dates(store, selected, force=args.force)
    logger.info("days: %d selected, %d already in store, %d to build",
                len(selected), len(selected) - len(pending), len(pending))
    if not pending:
        logger.info("nothing to do — store already warm for the selected range.")
        return 0

    params = dict(
        interpolation=args.interpolation, convention=DEFAULT_CONVENTION,
        calendar=DEFAULT_CALENDAR, modifier=DEFAULT_MODIFIER, spot_lag=DEFAULT_SPOT_LAG,
    )
    tasks = [(d, str(work_dir / f"{d.isoformat()}.parquet"), params) for d in pending]

    progress = Progress(total_days=len(tasks))
    logger.info("building %d days with %d workers (max_tasks_per_child=%s)...",
                len(tasks), args.n_jobs, args.max_tasks_per_child or "off")

    heartbeat_every = max(1, len(tasks) // 200)
    # max_tasks_per_child is OFF by default: worker memory is flat (~230 MB), and a
    # synchronized recycle wave (all workers hitting the limit at once) can deadlock
    # the pool on Windows spawn. Only recycle if the user explicitly opts in.
    mtpc = args.max_tasks_per_child or None
    with ProcessPoolExecutor(
        max_workers=args.n_jobs,
        initializer=_init_worker,
        initargs=(base_dir,),
        max_tasks_per_child=mtpc,
    ) as pool:
        futures = {pool.submit(_warm_one_day, t): t[0] for t in tasks}
        for fut in as_completed(futures):
            stat = fut.result()
            progress.record(stat)
            if stat.error:
                logger.warning("day %s FAILED: %s", stat.date, stat.error)
            if progress.days_done % heartbeat_every == 0 or progress.days_done == len(tasks):
                logger.info(progress.heartbeat())

    logger.info("DONE: %s", progress.heartbeat())
    logger.info("total curves written: %s across %d days (%d errors)",
                f"{progress.curves:,}", progress.days_done, progress.errors)

    if args.verify:
        _verify(store, pending, logger)
    return 0 if progress.errors == 0 else 1


def _verify(store: Any, dates: list[datetime.date], logger: logging.Logger) -> None:
    """Read back a sample day and confirm reconstructed curves have sane DFs."""
    if not dates:
        return
    sample = dates[len(dates) // 2]
    raw = store.read_raw_day(ASSET_NAME, sample)
    if raw is None or raw.empty:
        logger.warning("verify: no rows read back for %s", sample)
        return
    curves = store.reconstruct_curves_batch(raw.head(3), cfg=None, max_workers=2)
    ok = 0
    for ts, curve in curves.items():
        try:
            raw_nodes = curve.nodes._nodes if hasattr(curve.nodes, "_nodes") else dict(curve.nodes)
            dfs = [float(raw_nodes[d]) for d in sorted(raw_nodes.keys())]
            monotone = all(dfs[i] >= dfs[i + 1] - 1e-9 for i in range(len(dfs) - 1))
            if len(dfs) >= MIN_TENORS and abs(dfs[0] - 1.0) < 1e-6 and 0.0 < dfs[-1] < 1.0 and monotone:
                ok += 1
        except Exception as exc:
            logger.warning("verify check failed @ %s: %s", ts, exc)
    logger.info("verify: sample day %s -> %d rows, %d/%d reconstructed curves have sane monotone DFs",
                sample, len(raw), ok, len(curves))


def run_status(args: argparse.Namespace, logger: logging.Logger) -> int:
    from Caching.curve_store import CurveStore

    store = CurveStore(base_dir=Path(args.base_dir)) if args.base_dir else CurveStore.default()
    raw_dir = Path(store._base_dir) / "raw" / f"asset={ASSET_NAME}"
    parts = sorted(raw_dir.glob("date=*")) if raw_dir.exists() else []
    logger.info("asset=%s base_dir=%s", ASSET_NAME, store._base_dir)
    logger.info("populated day-partitions: %d", len(parts))
    if parts:
        logger.info("range: %s .. %s", parts[0].name, parts[-1].name)
        sample = parts[len(parts) // 2]
        d = datetime.date.fromisoformat(sample.name.split("=", 1)[1])
        raw = store.read_raw_day(ASSET_NAME, d)
        logger.info("sample %s: %d snapshot rows", d, 0 if raw is None else len(raw))
    return 0


# ─────────────────────────────────────────────────────────────────────────────
# CLI
# ─────────────────────────────────────────────────────────────────────────────
def _build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = p.add_subparsers(dest="command", required=True)

    default_jobs = max(1, min(28, (os.cpu_count() or 4) - 2))

    w = sub.add_parser("warm", help="Extract the workbook and warm the CurveStore.")
    w.add_argument("--workbook", default=DEFAULT_DB_PATH)
    w.add_argument("--work-dir", default=None, help="Dir for per-day par-rate parquet (default: ARBS cache).")
    w.add_argument("--base-dir", default=None, help="CurveStore base dir (default: ARBS cache / curve_store).")
    w.add_argument("--start-date", default=None, help="YYYY-MM-DD inclusive.")
    w.add_argument("--end-date", default=None, help="YYYY-MM-DD inclusive.")
    w.add_argument("--limit-days", type=int, default=None, help="Debug: build only the first N selected days.")
    w.add_argument("--n-jobs", type=int, default=default_jobs)
    w.add_argument("--max-tasks-per-child", type=int, default=None,
                   help="Recycle workers after N tasks (default off; a synchronized recycle wave can deadlock the pool).")
    w.add_argument("--interpolation", default=DEFAULT_INTERPOLATION)
    w.add_argument("--force", action="store_true", help="Rebuild days already present in the store.")
    w.add_argument("--force-extract", action="store_true", help="Re-extract day files even if present.")
    w.add_argument("--skip-extract", action="store_true", help="Reuse existing extracted day files.")
    w.add_argument("--verify", action="store_true", help="Read back a sample day after warming.")
    w.add_argument("--verbose", action="store_true")

    s = sub.add_parser("status", help="Report CurveStore population for the CitiVelo asset.")
    s.add_argument("--base-dir", default=None)
    s.add_argument("--verbose", action="store_true")

    return p


def main(argv: Optional[list[str]] = None) -> int:
    args = _build_parser().parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if getattr(args, "verbose", False) else logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        stream=sys.stdout,
    )
    if args.command == "warm":
        return run_warm(args, LOGGER)
    if args.command == "status":
        return run_status(args, LOGGER)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
