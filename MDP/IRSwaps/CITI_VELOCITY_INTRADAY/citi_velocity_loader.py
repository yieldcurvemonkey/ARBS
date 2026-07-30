"""Parse Citi Velocity ``CVTSHIST`` intraday SOFR-OIS par-rate workbooks.

Workbook layout (one sheet per week, named ``<startYYYYMMDDHHMM>-<endYYYYMMDDHHMM>``)::

    row 1: <start timestamp int>                 e.g. 202607130001
    row 2: <end   timestamp int>                 e.g. 202607171159
    row 3: (blank)
    row 4: =CVTSHIST("RATES.OIS.USD_SOFR.PAR.1D,...","MI01",,"...","...","CLOSE")
    row 5: Date | RATES.OIS.USD_SOFR.PAR.1D - CLOSE | ... (44 tenor columns)
    row 6+: <datetime> | <par rate %> | ...        (1-minute snapshots, newest first)

Only a subset of the weekly sheets are actually populated with data; the rest
are empty templates carrying just the two period-boundary timestamps.  The
parser detects the header row rather than hard-coding a row index, so it copes
with both the populated (``db.xlsx``) and the raw single-sheet
(``test.xlsx``) variants of the export.
"""

from __future__ import annotations

import datetime
import re
from typing import Dict, List, Optional, Sequence, Tuple

import openpyxl
import pandas as pd

__all__ = [
    "CURVE_TENOR_ORDER",
    "TENOR_HEADER_RE",
    "parse_tenor_from_header",
    "parse_period_from_sheet_name",
    "parse_sheet",
    "load_intraday_par_rates",
    "CitiVelocityWorkbook",
]

# Canonical curve tenor order as exported by Citi Velocity (short → long).
CURVE_TENOR_ORDER: Tuple[str, ...] = (
    "1D", "1W", "2W", "3W",
    "1M", "2M", "3M", "4M", "5M", "6M", "7M", "8M", "9M", "10M", "11M",
    "1Y", "15M", "18M", "21M",
    "2Y", "3Y", "4Y", "5Y", "6Y", "7Y", "8Y", "9Y", "10Y",
    "11Y", "12Y", "13Y", "14Y", "15Y", "16Y", "17Y", "18Y", "19Y", "20Y",
    "25Y", "30Y", "35Y", "40Y", "45Y", "50Y",
)

# e.g. "RATES.OIS.USD_SOFR.PAR.5Y - CLOSE" -> "5Y"
TENOR_HEADER_RE = re.compile(r"PAR\.([0-9]+[DWMY])\b", re.IGNORECASE)

_SHEET_NAME_RE = re.compile(r"^(\d{12})-(\d{12})$")


def _require_xlsx(path: str) -> None:
    """Reject legacy ``.xls`` (BIFF) paths with a consistent, actionable error.

    Called by every entry point that opens the workbook so metadata methods and
    data methods surface the same guidance instead of a raw openpyxl exception.
    """
    if not str(path).lower().endswith((".xlsx", ".xlsm")):
        raise ValueError(
            f"Citi Velocity loader expects an .xlsx workbook, got {path!r}. "
            "Legacy .xls (BIFF) exports must be converted to .xlsx first "
            "(openpyxl cannot read the OLE2 format)."
        )


def parse_tenor_from_header(header: object) -> Optional[str]:
    """Extract the rateslib tenor (e.g. ``"5Y"``) from a Citi column header."""
    if header is None:
        return None
    m = TENOR_HEADER_RE.search(str(header))
    return m.group(1).upper() if m else None


def _parse_ts_int(value: object) -> Optional[datetime.datetime]:
    """Parse a ``YYYYMMDDHHMM`` integer/string into a naive datetime."""
    if value is None:
        return None
    raw = str(int(value)) if isinstance(value, (int, float)) else str(value).strip()
    if len(raw) != 12 or not raw.isdigit():
        return None
    return datetime.datetime.strptime(raw, "%Y%m%d%H%M")


def parse_period_from_sheet_name(name: str) -> Optional[Tuple[datetime.datetime, datetime.datetime]]:
    """Return ``(start, end)`` datetimes encoded in a weekly sheet name.

    Returns ``None`` when the name does not match the ``YYYYMMDDHHMM-YYYYMMDDHHMM``
    convention.
    """
    m = _SHEET_NAME_RE.match(str(name).strip())
    if not m:
        return None
    start = _parse_ts_int(m.group(1))
    end = _parse_ts_int(m.group(2))
    if start is None or end is None:
        return None
    return start, end


def _find_header_row(rows: Sequence[Sequence[object]], max_scan: int = 12) -> Optional[int]:
    """Index (0-based within ``rows``) of the header row whose first cell is 'Date'."""
    for i, row in enumerate(rows[:max_scan]):
        if row and isinstance(row[0], str) and row[0].strip().lower() == "date":
            return i
    return None


def parse_sheet(
    ws: "openpyxl.worksheet.worksheet.Worksheet",
    *,
    tz: Optional[str] = None,
    restrict_to_known_tenors: bool = True,
) -> pd.DataFrame:
    """Parse one worksheet into a tidy time-indexed par-rate frame.

    Parameters
    ----------
    ws
        An openpyxl worksheet (read-only or normal).
    tz
        Optional IANA timezone to localize the (naive) source timestamps to.
        ``None`` (default) keeps them naive, exactly as stored.
    restrict_to_known_tenors
        When ``True`` unknown/extra columns are dropped and the surviving
        columns are ordered per :data:`CURVE_TENOR_ORDER`.

    Returns
    -------
    pandas.DataFrame
        Index = snapshot timestamp (sorted ascending, unique), columns = tenor
        strings, values = par swap rate in **percent**.  Empty template sheets
        yield an empty frame.
    """
    rows = list(ws.iter_rows(values_only=True))
    header_idx = _find_header_row(rows)
    if header_idx is None:
        return pd.DataFrame()

    header = rows[header_idx]
    # Map data columns -> tenor.  Column 0 is the timestamp.
    col_tenor: List[Tuple[int, str]] = []
    for j, cell in enumerate(header[1:], start=1):
        tenor = parse_tenor_from_header(cell)
        if tenor is not None:
            col_tenor.append((j, tenor))

    if not col_tenor:
        return pd.DataFrame()

    tenors = [t for _, t in col_tenor]
    records: List[Tuple[datetime.datetime, List[Optional[float]]]] = []
    for row in rows[header_idx + 1:]:
        if not row:
            continue
        ts = row[0]
        if not isinstance(ts, datetime.datetime):
            # data rows always start with a datetime; anything else is padding
            continue
        values: List[Optional[float]] = []
        any_value = False
        for j, _ in col_tenor:
            v = row[j] if j < len(row) else None
            if v is None or (isinstance(v, str) and not v.strip()):
                values.append(None)
                continue
            try:
                values.append(float(v))
                any_value = True
            except (TypeError, ValueError):
                values.append(None)
        if any_value:
            records.append((ts, values))

    if not records:
        return pd.DataFrame()

    index = pd.DatetimeIndex([r[0] for r in records], name="timestamp")
    df = pd.DataFrame([r[1] for r in records], index=index, columns=tenors, dtype="float64")

    # newest-first in the source; present ascending and de-duplicate.
    df = df[~df.index.duplicated(keep="first")].sort_index()

    if restrict_to_known_tenors:
        ordered = [t for t in CURVE_TENOR_ORDER if t in df.columns]
        df = df[ordered]

    if tz is not None:
        df.index = df.index.tz_localize(tz)

    return df


def _iter_populated_sheet_names(wb: "openpyxl.workbook.workbook.Workbook") -> List[str]:
    out = []
    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.max_column and ws.max_column > 1:
            out.append(sn)
    return out


def load_intraday_par_rates(
    path: str,
    *,
    sheets: Optional[Sequence[str]] = None,
    tz: Optional[str] = None,
    concat: bool = True,
) -> pd.DataFrame | Dict[str, pd.DataFrame]:
    """Load Citi Velocity intraday par rates from an ``.xlsx`` workbook.

    Parameters
    ----------
    path
        Path to the ``.xlsx`` workbook (``db.xlsx`` / ``test.xlsx``).
    sheets
        Explicit sheet names to parse.  ``None`` (default) auto-detects every
        populated sheet (i.e. skips the empty weekly templates).
    tz
        Optional timezone to localize timestamps to (see :func:`parse_sheet`).
    concat
        When ``True`` (default) return a single time-indexed frame spanning all
        parsed sheets; otherwise return ``{sheet_name: frame}``.

    Returns
    -------
    Either a single :class:`pandas.DataFrame` (``concat=True``) or a mapping of
    sheet name to frame.
    """
    _require_xlsx(path)

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(sheets) if sheets is not None else _iter_populated_sheet_names(wb)
        frames: Dict[str, pd.DataFrame] = {}
        for sn in names:
            frame = parse_sheet(wb[sn], tz=tz)
            if not frame.empty:
                frames[sn] = frame
    finally:
        wb.close()

    if not concat:
        return frames

    if not frames:
        return pd.DataFrame()

    combined = pd.concat(frames.values(), axis=0)
    combined = combined[~combined.index.duplicated(keep="last")].sort_index()
    return combined


class CitiVelocityWorkbook:
    """Lazy, cached reader over a Citi Velocity intraday workbook.

    >>> wb = CitiVelocityWorkbook("db.xlsx")
    >>> wb.populated_sheets()          # weekly tabs that carry data
    >>> frame = wb.frame()             # all snapshots, one tidy time-indexed frame
    >>> series = wb.snapshot(ts)       # the par curve nearest a timestamp
    """

    def __init__(self, path: str, *, tz: Optional[str] = None):
        self.path = str(path)
        self.tz = tz
        self._frames: Optional[Dict[str, pd.DataFrame]] = None
        self._combined: Optional[pd.DataFrame] = None
        self._sheet_names: Optional[List[str]] = None

    # -- sheet metadata -------------------------------------------------
    def sheet_names(self) -> List[str]:
        if self._sheet_names is None:
            _require_xlsx(self.path)
            wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
            try:
                self._sheet_names = list(wb.sheetnames)
            finally:
                wb.close()
        return list(self._sheet_names)

    def populated_sheets(self) -> List[str]:
        return list(self.frames().keys())

    def sheet_periods(self) -> Dict[str, Tuple[datetime.datetime, datetime.datetime]]:
        out = {}
        for sn in self.sheet_names():
            period = parse_period_from_sheet_name(sn)
            if period is not None:
                out[sn] = period
        return out

    # -- data -----------------------------------------------------------
    def frames(self) -> Dict[str, pd.DataFrame]:
        if self._frames is None:
            self._frames = load_intraday_par_rates(self.path, tz=self.tz, concat=False)  # type: ignore[assignment]
        return self._frames

    def frame(self) -> pd.DataFrame:
        """All populated snapshots as one ascending, de-duplicated frame."""
        if self._combined is None:
            frames = self.frames()
            if not frames:
                self._combined = pd.DataFrame()
            else:
                combined = pd.concat(frames.values(), axis=0)
                self._combined = combined[~combined.index.duplicated(keep="last")].sort_index()
        return self._combined

    def timestamps(self) -> pd.DatetimeIndex:
        return self.frame().index

    def snapshot(
        self,
        when: Optional[datetime.datetime] = None,
        *,
        method: str = "nearest",
    ) -> pd.Series:
        """Return the par-rate curve at (or nearest/asof) ``when``.

        ``method`` is one of ``"nearest"``, ``"asof"`` (<= when) or ``"exact"``.
        ``when=None`` returns the most recent snapshot.
        """
        frame = self.frame()
        if frame.empty:
            raise ValueError(f"No populated snapshots in workbook {self.path!r}")
        if when is None:
            return frame.iloc[-1].dropna()

        idx = frame.index
        key = pd.Timestamp(when)
        if key.tzinfo is None and idx.tz is not None:
            key = key.tz_localize(idx.tz)
        elif key.tzinfo is not None and idx.tz is None:
            key = key.tz_localize(None)

        if method == "exact":
            row = frame.loc[key]
        elif method == "asof":
            pos = idx.get_indexer([key], method="ffill")[0]
            if pos == -1:
                raise KeyError(f"No snapshot at or before {when!r}")
            row = frame.iloc[pos]
        elif method == "nearest":
            pos = idx.get_indexer([key], method="nearest")[0]
            row = frame.iloc[pos]
        else:
            raise ValueError(f"Unknown method {method!r}; use nearest|asof|exact")
        return row.dropna()
